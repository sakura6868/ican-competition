"""手册/文档解析：PDF、DOCX、纯文本 → 文本内容。"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path


class ParseError(Exception):
    pass


@dataclass(frozen=True)
class ParsedSection:
    text: str
    page: int | None = None
    section: str | None = None


def extract_text(path: Path) -> str:
    return "\n".join(section.text for section in extract_sections(path))


def extract_sections(path: Path) -> list[ParsedSection]:
    """提取带真实页码/标题的文本段；无页概念时不伪造页码。"""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf_sections(path)
    if suffix in (".docx", ".doc"):
        return _extract_docx_sections(path)
    if suffix in (".txt", ".md", ".csv"):
        return _extract_plain_sections(path)
    raise ParseError(f"不支持的文档类型: {suffix}")


def _extract_pdf_sections(path: Path) -> list[ParsedSection]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise ParseError("缺少 pdfplumber（uv sync --extra media）") from exc
    pages: list[ParsedSection] = []
    with pdfplumber.open(path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            stripped = text.strip()
            if not stripped:
                continue
            first_line = next((line.strip() for line in stripped.splitlines() if line.strip()), "")
            section = first_line if 0 < len(first_line) <= 160 else None
            pages.append(ParsedSection(text=stripped, page=page_number, section=section))
    return pages


def _extract_docx_sections(path: Path) -> list[ParsedSection]:
    try:
        import docx
    except ImportError as exc:
        raise ParseError("缺少 python-docx（uv sync --extra media）") from exc
    document = docx.Document(str(path))
    sections: list[ParsedSection] = []
    current_title: str | None = None
    parts: list[str] = []

    def flush() -> None:
        if parts:
            sections.append(ParsedSection(text="\n".join(parts), section=current_title))
            parts.clear()

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = str(getattr(para.style, "name", "") or "").lower()
        if style_name.startswith("heading"):
            flush()
            current_title = text
            parts.append(text)
        else:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    flush()
    return sections


def _extract_plain_sections(path: Path) -> list[ParsedSection]:
    text = path.read_text(encoding="utf-8", errors="replace")
    if path.suffix.lower() != ".md":
        stripped = text.strip()
        return [ParsedSection(text=stripped)] if stripped else []

    sections: list[ParsedSection] = []
    current_title: str | None = None
    parts: list[str] = []

    def flush() -> None:
        if parts:
            body = "\n".join(parts).strip()
            if body:
                sections.append(ParsedSection(text=body, section=current_title))
            parts.clear()

    for line in text.splitlines():
        match = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*$", line)
        if match:
            flush()
            current_title = match.group(1).strip()
        parts.append(line)
    flush()
    return sections
